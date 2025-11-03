"""Core resume generation logic using python-docx."""

import json
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)


class ResumeGenerator:
    """Generates ATS-friendly resumes in DOCX format from JSON input."""
    
    def __init__(self, json_data: Dict[str, Any], config_path: Optional[Path] = None):
        """Initialize generator with JSON data and optional config.
        
        Args:
            json_data: Dictionary containing resume data
            config_path: Optional path to config.json file
        """
        self.data = json_data
        self.config = self._load_config(config_path)
        self.doc = Document()
        self._setup_document()
    
    def _load_config(self, config_path: Optional[Path] = None) -> Dict[str, Any]:
        """Load configuration from config.json file.
        
        Args:
            config_path: Optional path to config file
            
        Returns:
            Dictionary containing configuration
        """
        if config_path is None:
            # Default to config.json in same directory as this file
            config_path = Path(__file__).parent / "config.json"
        
        try:
            with open(config_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except FileNotFoundError:
            # Return default configuration if file not found
            logger.warning(f"Config file not found at {config_path}, using default configuration")
            return self._get_default_config()
    
    def _get_default_config(self) -> Dict[str, Any]:
        """Return default configuration if config.json is not found.
        
        Returns:
            Dictionary with default configuration values
        """
        return {
            "fonts": {"primary": "Calibri", "fallback": "Arial"},
            "font_sizes": {"body": 11, "name": 16, "section_header": 12},
            "spacing": {
                "margin": 36,
                "header_name_after": 2,
                "header_contact_after": 8,
                "section_header_after": 4,
                "bullet_after": 1,
                "subsection_header_after": 1,
                "inter_subsection_spacing": 0,
                "inter_job_spacing": 4,
                "skill_category_after": 2,
                "education_detail_after": 4,
                "professional_summary_paragraph": 4,
                "professional_summary_last": 6
            },
            "colors": {
                "text": {"r": 0, "g": 0, "b": 0},
                "hyperlink": "0563C1"
            },
            "formatting": {
                "underline_name": True,
                "underline_section_headers": True,
                "underline_subsection_headers": True,
                "italic_skills": True,
                "italic_job_title": True,
                "contact_separator": " • "
            }
        }
    
    @property
    def FONT_NAME(self) -> str:
        """Get primary font name from config."""
        return self.config["fonts"]["primary"]
    
    @property
    def BODY_FONT_SIZE(self) -> int:
        """Get body font size from config."""
        return self.config["font_sizes"]["body"]
    
    @property
    def NAME_FONT_SIZE(self) -> int:
        """Get name font size from config."""
        return self.config["font_sizes"]["name"]
    
    @property
    def SECTION_HEADER_FONT_SIZE(self) -> int:
        """Get section header font size from config."""
        return self.config["font_sizes"]["section_header"]
    
    @property
    def MARGIN_SIZE(self) -> int:
        """Get margin size from config."""
        return self.config["spacing"]["margin"]
    
    def _setup_document(self):
        """Setup document margins and default styles."""
        # Set narrow margins for better space usage (0.5 inches)
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Pt(self.MARGIN_SIZE)
            section.bottom_margin = Pt(self.MARGIN_SIZE)
            section.left_margin = Pt(self.MARGIN_SIZE)
            section.right_margin = Pt(self.MARGIN_SIZE)
        
        # Modify List Bullet style to have minimal spacing
        try:
            list_bullet_style = self.doc.styles['List Bullet']
            list_bullet_style.paragraph_format.space_after = Pt(0)
            list_bullet_style.paragraph_format.space_before = Pt(0)
        except KeyError:
            # Style doesn't exist yet, will be created when first bullet is added
            pass
    
    def _add_paragraph(self, text: str, bold: bool = False, 
                      font_size: int = None, alignment: str = "left",
                      space_after: int = 6, italic: bool = False,
                      underline: bool = False) -> None:
        """Add a formatted paragraph to the document.
        
        Args:
            text: Text content
            bold: Whether text should be bold
            font_size: Font size in points (defaults to BODY_FONT_SIZE)
            alignment: Text alignment (left, center, right)
            space_after: Space after paragraph in points
            italic: Whether text should be italic
            underline: Whether text should be underlined
        """
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        
        # Set font properties
        run.font.name = self.FONT_NAME
        run.font.size = Pt(font_size if font_size else self.BODY_FONT_SIZE)
        run.bold = bold
        run.italic = italic
        run.underline = underline
        text_color = self.config["colors"]["text"]
        run.font.color.rgb = RGBColor(text_color["r"], text_color["g"], text_color["b"])
        
        # Set alignment
        if alignment == "center":
            para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        elif alignment == "right":
            para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        else:
            para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        
        # Set spacing
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(0)
    
    def _add_hyperlink(self, paragraph, text: str, url: str):
        """Add a clickable hyperlink to a paragraph.
        
        Args:
            paragraph: The paragraph to add the hyperlink to
            text: Display text for the link
            url: URL to link to
        """
        # Get the document part and create a relationship
        part = paragraph.part
        r_id = part.relate_to(
            url, 
            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
            is_external=True
        )
        
        # Create the hyperlink element
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), r_id)
        
        # Create a new run element
        new_run = OxmlElement('w:r')
        
        # Create run properties
        rPr = OxmlElement('w:rPr')
        
        # Set font properties
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), self.FONT_NAME)
        rFonts.set(qn('w:hAnsi'), self.FONT_NAME)
        rPr.append(rFonts)
        
        # Set font size
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(self.BODY_FONT_SIZE * 2))  # Half-points
        rPr.append(sz)
        
        # Set color to blue for hyperlink
        color = OxmlElement('w:color')
        hyperlink_color = self.config["colors"]["hyperlink"]
        color.set(qn('w:val'), hyperlink_color)
        rPr.append(color)
        
        # Set underline
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
        
        new_run.append(rPr)
        
        # Add text element
        text_elem = OxmlElement('w:t')
        text_elem.text = text
        new_run.append(text_elem)
        
        hyperlink.append(new_run)
        paragraph._element.append(hyperlink)

    def _add_header_section(self):
        """Add header section with name and contact information."""
        header_data = self.data.get("header", {})
        
        # Name - 16pt bold and underlined
        name = header_data.get("name", "")
        if name:
            underline = self.config["formatting"]["underline_name"]
            spacing = self.config["spacing"]["header_name_after"]
            self._add_paragraph(name, bold=True, font_size=self.NAME_FONT_SIZE, 
                              alignment="center", space_after=spacing, underline=underline)
        
        # Contact info - single line, centered with clickable links
        para = self.doc.add_paragraph()
        para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        para.paragraph_format.space_after = Pt(self.config["spacing"]["header_contact_after"])
        para.paragraph_format.space_before = Pt(0)
        
        # Build contact line with hyperlinks
        contact_items = []
        
        # Email (clickable)
        if "email" in header_data:
            email = header_data["email"]
            if email:  # Check if not None
                contact_items.append(("email", email, f"mailto:{email}"))
        
        # Phone (not clickable)
        if "phone" in header_data:
            phone = header_data["phone"]
            if phone:  # Check if not None
                contact_items.append(("text", phone, None))
        
        # Location (not clickable)
        if "location" in header_data:
            location = header_data["location"]
            if location:  # Check if not None
                contact_items.append(("text", location, None))
        
        # LinkedIn (clickable)
        if "linkedin" in header_data:
            linkedin = header_data["linkedin"]
            if linkedin:  # Check if not None
                # Add https:// if not present
                url = linkedin if linkedin.startswith("http") else f"https://{linkedin}"
                contact_items.append(("link", linkedin, url))
        
        # GitHub (clickable)
        if "github" in header_data:
            github = header_data["github"]
            if github:  # Check if not None
                # Add https:// if not present
                url = github if github.startswith("http") else f"https://{github}"
                contact_items.append(("link", github, url))
        
        # Add all contact items with separators
        separator = self.config["formatting"]["contact_separator"]
        for i, (item_type, text, url) in enumerate(contact_items):
            if i > 0:
                # Add separator
                run = para.add_run(separator)
                run.font.name = self.FONT_NAME
                run.font.size = Pt(self.BODY_FONT_SIZE)
                text_color = self.config["colors"]["text"]
                run.font.color.rgb = RGBColor(text_color["r"], text_color["g"], text_color["b"])
            
            if item_type in ["email", "link"]:
                # Add as hyperlink
                self._add_hyperlink(para, text, url)
            else:
                # Add as regular text
                run = para.add_run(text)
                run.font.name = self.FONT_NAME
                run.font.size = Pt(self.BODY_FONT_SIZE)
                text_color = self.config["colors"]["text"]
                run.font.color.rgb = RGBColor(text_color["r"], text_color["g"], text_color["b"])
    
    def _add_section_header(self, title: str):
        """Add a section header (uppercase, 12pt bold, underlined).
        
        Args:
            title: Section title
        """
        underline = self.config["formatting"]["underline_section_headers"]
        spacing = self.config["spacing"]["section_header_after"]
        self._add_paragraph(title.upper(), bold=True, 
                          font_size=self.SECTION_HEADER_FONT_SIZE,
                          space_after=spacing, underline=underline)
    
    def _add_professional_summary_section(self):
        """Add professional summary section (optional).
        
        This section provides a brief overview of professional background
        and appears before skills and experience sections.
        """
        summary_data = self.data.get("professional_summary")
        if not summary_data:
            return
        
        self._add_section_header("Professional Summary")
        
        # Summary can be either a string or a list of paragraphs
        if isinstance(summary_data, str):
            spacing = self.config["spacing"]["professional_summary_last"]
            self._add_paragraph(summary_data, space_after=spacing)
        elif isinstance(summary_data, list):
            for i, paragraph in enumerate(summary_data):
                # Add spacing after last paragraph
                if i == len(summary_data) - 1:
                    spacing = self.config["spacing"]["professional_summary_last"]
                else:
                    spacing = self.config["spacing"]["professional_summary_paragraph"]
                self._add_paragraph(paragraph, space_after=spacing)
    
    def _add_bullet_paragraph(self, text: str):
        """Add a bullet point paragraph with proper formatting.
        
        Args:
            text: Bullet point text
        """
        para = self.doc.add_paragraph(text, style='List Bullet')
        spacing = self.config["spacing"]["bullet_after"]
        
        # Set paragraph spacing - both space_after AND line_spacing
        para.paragraph_format.space_after = Pt(spacing)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing = 1.0  # Single line spacing
        
        # Set font for bullet text
        for run in para.runs:
            run.font.name = self.FONT_NAME
            run.font.size = Pt(self.BODY_FONT_SIZE)
            text_color = self.config["colors"]["text"]
            run.font.color.rgb = RGBColor(text_color["r"], text_color["g"], text_color["b"])
    
    def _add_skills_section(self):
        """Add skills section."""
        skills_data = self.data.get("skills", [])
        if not skills_data:
            return
        
        self._add_section_header("Skills")
        
        # Group skills by category if provided, otherwise show as list
        if isinstance(skills_data, dict):
            for category, skills in skills_data.items():
                if isinstance(skills, list):
                    skills_text = ", ".join(skills)
                else:
                    skills_text = skills
                italic = self.config["formatting"]["italic_skills"]
                spacing = self.config["spacing"]["skill_category_after"]
                self._add_paragraph(f"{category}: {skills_text}", italic=italic, space_after=spacing)
        elif isinstance(skills_data, list):
            skills_text = ", ".join(skills_data)
            self._add_paragraph(skills_text, space_after=6)
    
    def _add_experience_section(self):
        """Add experience section with roles, companies, dates, and bullets.
        
        Supports both traditional bullets and subsection-based organization.
        """
        experience_data = self.data.get("experience", [])
        if not experience_data:
            return
        
        self._add_section_header("Relevant Experience")
        
        for i, job in enumerate(experience_data):
            # Company and location on left, dates on right (same line)
            company = job.get("company", "")
            location = job.get("location", "")
            dates = job.get("dates", "")
            
            # Create a paragraph with company/location on left and dates on right
            para = self.doc.add_paragraph()
            
            # Left side: company and location
            left_parts = []
            if company:
                left_parts.append(company)
            if location:
                left_parts.append(location)
            
            if left_parts:
                left_text = ", ".join(left_parts)
                run_left = para.add_run(left_text)
                run_left.font.name = self.FONT_NAME
                run_left.font.size = Pt(self.BODY_FONT_SIZE)
                run_left.bold = True
                text_color = self.config["colors"]["text"]
                run_left.font.color.rgb = RGBColor(text_color["r"], text_color["g"], text_color["b"])
            
            # Add tab to push dates to the right
            if dates and left_parts:
                # Add tab
                para.add_run("\t")
                
                # Right side: dates
                run_right = para.add_run(dates)
                run_right.font.name = self.FONT_NAME
                run_right.font.size = Pt(self.BODY_FONT_SIZE)
                run_right.bold = True
                text_color = self.config["colors"]["text"]
                run_right.font.color.rgb = RGBColor(text_color["r"], text_color["g"], text_color["b"])
                
                # Set right-aligned tab stop
                tab_stops = para.paragraph_format.tab_stops
                tab_stops.add_tab_stop(Inches(6.0), alignment=WD_TAB_ALIGNMENT.RIGHT)
            
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.space_before = Pt(0)
            
            # Job title - italic on its own line
            role = job.get("role", "")
            if role:
                italic = self.config["formatting"]["italic_job_title"]
                self._add_paragraph(role, italic=italic, space_after=2)
            
            # Check for subsections (new feature) or traditional bullets
            subsections = job.get("subsections", [])
            bullets = job.get("bullets", [])
            
            if subsections:
                # Use subsection-based organization
                for subsection in subsections:
                    subsection_header = subsection.get("header", "")
                    subsection_bullets = subsection.get("bullets", [])
                    
                    # Add subsection header (uppercase, underlined)
                    if subsection_header:
                        underline = self.config["formatting"]["underline_subsection_headers"]
                        spacing = self.config["spacing"]["subsection_header_after"]
                        self._add_paragraph(
                            subsection_header.upper() + ":",
                            bold=False,
                            underline=underline,
                            space_after=spacing,
                            font_size=self.BODY_FONT_SIZE
                        )
                    
                    # Add bullets under this subsection
                    for bullet in subsection_bullets:
                        self._add_bullet_paragraph(bullet)
                    
                    # Add spacing between subsections (only if configured and not last)
                    is_last = subsection == subsections[-1]
                    if not is_last:
                        spacing = self.config["spacing"]["inter_subsection_spacing"]
                        # Only add empty paragraph if spacing > 0 to avoid blank lines
                        if spacing > 0:
                            self._add_paragraph("", space_after=spacing)
            else:
                # Traditional bullet format (backward compatible)
                for bullet in bullets:
                    self._add_bullet_paragraph(bullet)
            
            # Add space after each job (except the last one)
            if i < len(experience_data) - 1:
                spacing = self.config["spacing"]["inter_job_spacing"]
                self._add_paragraph("", space_after=spacing)
    
    def _add_education_section(self):
        """Add education section."""
        education_data = self.data.get("education", [])
        if not education_data:
            return
        
        self._add_section_header("Education")
        
        for edu in education_data:
            # Degree and institution - bold
            degree = edu.get("degree", "")
            institution = edu.get("institution", "")
            title_line = f"{degree}, {institution}" if degree and institution else degree or institution
            
            if title_line:
                self._add_paragraph(title_line, bold=True, space_after=1)
            
            # Dates and location
            dates = edu.get("dates", "")
            location = edu.get("location", "")
            gpa = edu.get("gpa", "")
            
            detail_parts = []
            if dates:
                detail_parts.append(dates)
            if location:
                detail_parts.append(location)
            if gpa:
                detail_parts.append(f"GPA: {gpa}")
            
            if detail_parts:
                detail_line = " | ".join(detail_parts)
                spacing = self.config["spacing"]["education_detail_after"]
                self._add_paragraph(detail_line, space_after=spacing)
    
    def _add_awards_section(self):
        """Add awards and achievements section."""
        awards_data = self.data.get("awards", [])
        if not awards_data:
            return
        
        self._add_section_header("Awards")
        
        for award in awards_data:
            if isinstance(award, str):
                # Simple string award
                self._add_bullet_paragraph(award)
            elif isinstance(award, dict):
                # Award with details
                title = award.get("title", "")
                date = award.get("date", "")
                description = award.get("description", "")
                
                award_text = title
                if date:
                    award_text += f" ({date})"
                if description:
                    award_text += f" - {description}"
                
                self._add_bullet_paragraph(award_text)
    
    def generate(self, output_path: Path):
        """Generate the resume document.
        
        Args:
            output_path: Path where the DOCX file will be saved
        """
        # Add all sections in order
        self._add_header_section()
        self._add_professional_summary_section()
        self._add_skills_section()
        self._add_experience_section()
        self._add_education_section()
        self._add_awards_section()
        
        # Save the document
        self.doc.save(str(output_path))


def generate_resume_from_json(json_path: Path, output_path: Path):
    """Generate resume from JSON file.
    
    Args:
        json_path: Path to input JSON file
        output_path: Path to output DOCX file
    """
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    generator = ResumeGenerator(data)
    generator.generate(output_path)
